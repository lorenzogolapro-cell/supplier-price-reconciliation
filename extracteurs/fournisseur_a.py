# -*- coding: utf-8 -*-
"""
Extracteur du catalogue tarifaire du FOURNISSEUR A (fichier Excel,
tarif_fournisseur.xlsx).

Particularite du fichier : les ~15 premieres lignes sont un en-tete de
courtoisie (adresse, conditions de port, franco, mention "TARIF DISTRIBUTEUR
CONFIDENTIEL"). La vraie ligne d'en-tete est detectee dynamiquement en
cherchant la premiere ligne contenant a la fois "Reference" et "Designation",
pour que le code resiste a un decalage dans les millesimes suivants.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from extracteurs.base import (
    calculer_prix_colis,
    chemin_lisible,
    controler_coherence_prix,
    economie_palier,
    finaliser,
    nettoyer_ean,
    nettoyer_nombre,
    nettoyer_texte,
    normaliser_conditionnement,
    resumer_paliers,
    trouver_ligne_entete,
    _normaliser,
)

FOURNISSEUR = "FOURNISSEUR A"

# Mots-cles servant a reperer la ligne d'en-tete reelle
MOTS_CLES_ENTETE = ["Référence", "Désignation"]

# Correspondance en-tete du fichier -> champ du schema normalise.
# La cle est l'en-tete normalisee (minuscules, sans accent ni ponctuation).
# On matche par NOM et non par position : c'est ce qui evite de confondre
# "Votre tarif net HT unitaire" avec les paliers "Votre tarif net HT 2/3".
ENTETES_VERS_CHAMPS = {
    "reference": "ref_fournisseur",
    "designation": "designation",
    # ATTENTION : dans le tarif 2026 cette colonne est intitulee
    # "Tarif public conseille TTC" -> c'est bien du TTC, verifie sur les
    # donnees (prix_achat = TTC / (1 + TVA) x (1 - remise)).
    "tarifpublicconseillettc": "tarif_public_ttc",
    "tva": "tva_taux",
    "remise": "remise_taux",
    "votretarifnethtunitaire": "prix_achat_unitaire_ht",
    # Prix degressifs : 878 lignes ont un palier 2, 468 un palier 3.
    # Ce sont des PRIX UNITAIRES applicables a partir de la quantite seuil,
    # pas des prix de lot.
    "qte2": "palier2_qte",
    "votretarifnetht2": "palier2_prix_ht",
    "qte3": "palier3_qte",
    "votretarifnetht3": "palier3_prix_ht",
    "codelppr": "code_lppr",
    "montantlppr": "montant_lppr",
    "ecopartht": "eco_part_ht",
    "codeean": "ean",
    "conditionnement": "conditionnement",
    "dm": "dispositif_medical",
    "origine": "origine",
}

# Filet de securite : positions 0-based attendues d'apres le tarif 2026.
# Utilisees uniquement pour les champs que le matching par nom n'a pas
# trouves (en-tete renomme, cellule fusionnee, etc.).
INDEX_SECOURS = {
    "ref_fournisseur": 1,
    "designation": 2,
    "tarif_public_ttc": 3,
    "tva_taux": 4,
    "remise_taux": 5,
    "prix_achat_unitaire_ht": 6,
    "palier2_qte": 7,
    "palier2_prix_ht": 8,
    "palier3_qte": 9,
    "palier3_prix_ht": 10,
    "code_lppr": 11,
    "montant_lppr": 12,
    "eco_part_ht": 14,
    "ean": 16,
    "conditionnement": 17,
    "dispositif_medical": 19,
    "origine": 21,
}

# Colonnes volontairement ignorees :
#   0  Page              -> mise en page du catalogue papier
#   13 Packaging retail
#   15 Eco Part TTC      -> recalculable depuis eco_part_ht
#   18 Poids brut
#   20 Code douanier

# Champs a convertir en nombre / en texte
CHAMPS_NUM = {
    "tarif_public_ttc",
    "tva_taux",
    "remise_taux",
    "prix_achat_unitaire_ht",
    "palier2_qte",
    "palier2_prix_ht",
    "palier3_qte",
    "palier3_prix_ht",
    "montant_lppr",
    "eco_part_ht",
}


def _construire_mapping(feuille, ligne_entete: int) -> dict[str, int]:
    """Associe chaque champ du schema a un index de colonne 0-based."""
    entetes = next(
        feuille.iter_rows(
            min_row=ligne_entete, max_row=ligne_entete, values_only=True
        )
    )

    mapping: dict[str, int] = {}
    for index, entete in enumerate(entetes):
        if entete is None:
            continue
        champ = ENTETES_VERS_CHAMPS.get(_normaliser(str(entete)))
        # Le premier trouve gagne : evite qu'un doublon d'en-tete ecrase
        # la bonne colonne.
        if champ and champ not in mapping:
            mapping[champ] = index

    # Complement par les positions connues pour les champs manquants
    for champ, index in INDEX_SECOURS.items():
        if champ not in mapping and index < len(entetes):
            mapping[champ] = index

    return mapping


def extract(path) -> "pd.DataFrame":  # noqa: F821 (type resolu a l'execution)
    """Extrait le catalogue du FOURNISSEUR A vers le schema normalise commun."""
    path = Path(path)
    # read_only : le fichier fait ~3100 lignes, on evite de tout charger en RAM.
    # data_only : on veut les valeurs calculees, pas les formules.
    classeur = openpyxl.load_workbook(
        chemin_lisible(path), read_only=True, data_only=True
    )
    feuille = classeur[classeur.sheetnames[0]]

    ligne_entete = trouver_ligne_entete(feuille, MOTS_CLES_ENTETE)
    mapping = _construire_mapping(feuille, ligne_entete)

    manquants = [c for c in ("ref_fournisseur", "designation") if c not in mapping]
    if manquants:
        raise ValueError(f"Colonnes indispensables introuvables : {manquants}")

    lignes: list[dict] = []
    for valeurs in feuille.iter_rows(min_row=ligne_entete + 1, values_only=True):
        if valeurs is None:
            continue

        def cellule(champ):
            """Valeur brute d'un champ, ou None si la colonne est absente."""
            index = mapping.get(champ)
            if index is None or index >= len(valeurs):
                return None
            return valeurs[index]

        ref = nettoyer_texte(cellule("ref_fournisseur"))
        designation = nettoyer_texte(cellule("designation"))

        # On ignore les lignes sans reference ou sans designation :
        # separateurs de rubrique, lignes vides, pieds de page.
        if not ref or not designation:
            continue

        ligne = {
            "fournisseur": FOURNISSEUR,
            "ref_fournisseur": ref,
            "designation": designation,
            "ean": nettoyer_ean(cellule("ean")),
            "code_lppr": nettoyer_texte(cellule("code_lppr")),
            "dispositif_medical": nettoyer_texte(cellule("dispositif_medical")),
            "origine": nettoyer_texte(cellule("origine")),
            "fichier_source": path.name,
        }
        for champ in CHAMPS_NUM:
            ligne[champ] = nettoyer_nombre(cellule(champ))

        # Conditionnement : absent ou nul -> 1, et on garde la trace.
        (
            ligne["conditionnement"],
            ligne["conditionnement_suppose"],
        ) = normaliser_conditionnement(cellule("conditionnement"))

        # Prix du colis complet = prix unitaire x nombre d'unites par colis.
        ligne["prix_colis_ht"] = calculer_prix_colis(
            ligne["prix_achat_unitaire_ht"], ligne["conditionnement"]
        )

        # Controle : le prix net du fichier correspond-il bien au tarif
        # public TTC diminue de la TVA puis de la remise ?
        (
            ligne["prix_recalcule"],
            ligne["ecart_prix"],
            ligne["prix_coherent"],
        ) = controler_coherence_prix(
            ligne["tarif_public_ttc"],
            ligne["tva_taux"],
            ligne["remise_taux"],
            ligne["prix_achat_unitaire_ht"],
        )

        # Paliers degressifs : prix unitaires applicables au-dela d'un seuil
        prix_base = ligne["prix_achat_unitaire_ht"]
        ligne["economie_palier2_pct"] = economie_palier(
            prix_base, ligne["palier2_prix_ht"]
        )
        ligne["economie_palier3_pct"] = economie_palier(
            prix_base, ligne["palier3_prix_ht"]
        )
        ligne["paliers"] = resumer_paliers(
            prix_base,
            [
                (ligne["palier2_qte"], ligne["palier2_prix_ht"]),
                (ligne["palier3_qte"], ligne["palier3_prix_ht"]),
            ],
        )

        lignes.append(ligne)

    classeur.close()
    return finaliser(lignes)
